# -*- coding: utf-8 -*-
"""다국어 글자 표 모으기 → Kachi_Gun/Table/String.xlsx (칸: Key = 한국어 · EN = 영어 · Where = 나오는 곳)
사용: python tools/strings_collect.py
 - 코드(Assets/Scripts, Editor 제외)의 한글 글자 + 표에서 화면에 나오는 이름·설명을 모음
 - 이미 있는 줄의 EN 은 그대로 둠 (직접 고친 영어는 안 지워짐). 새 줄은 아래 EN 사전으로 채우고, 없으면 비워 둠(게임에서는 한국어로 보임)
 - 표 안 줄바꿈은 \\n 두 글자로 적음 (게임이 줄바꿈으로 바꿈)
 - 마지막에 csv 도 만듦 (xlsx2csv)
"""
import csv, glob, io, os, re, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

TOOLS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, TOOLS)
ROOT = os.path.join(TOOLS, '..', 'Kachi_Gun')
SCRIPTS = os.path.join(ROOT, 'Assets', 'Scripts')
TABLES = os.path.join(ROOT, 'Assets', 'Resources', 'Tables')
OUT = os.path.join(ROOT, 'Table', 'String.xlsx')
KO = re.compile('[가-힣]')

# 표에서 화면에 나오는 칸
TABLE_COLS = [
    ('Skin_Table', ['Name']), ('Weapon_Master', ['Name']), ('Zombie_Master', ['Name']),
    ('Supply_Collection', ['Name', 'Grade']), ('Mission_Master', ['Mission_Name']),
    ('Shop_Products', ['Name', 'Description']), ('Chapter_Summary', ['Biome_Name']),
    ('Unlock_Roadmap', ['Lock_Toast_Message', 'Unlock_Banner_Message']), ('Mine_Ore_Spec', ['Name_KR']),
]
# 코드에 있어도 번역하지 않는 것 (출시 전에 지울 테스트 버튼 등)
SKIP = {'[테스트용 — 에디터에서만 보임]', '다이아 +1000', '골드 x100', '다음 스테이지', '처음부터', '세이브를 지우고 처음부터 시작할까요?',
        '100 다이아', '언어 · Language', '한국어'}

EN = {
    # ---------- 전투 화면 ----------
    '절전모드': 'Eco', '상점': 'Shop', '미션': 'Missions', '자동강화': 'Auto', '공격': 'ATK', '골드': 'Gold', '다이아': 'Gems',
    '보스 도전': 'Fight Boss', '도전': 'Challenge', '도전 {0}': 'Challenge {0}', '공격력': 'Attack', '공격속도': 'Attack Speed', '돈 획득량': 'Gold Gain',
    '총알 {0}': 'Ammo {0}', '비어 있음': 'Empty', '무기': 'Weapons', '스킨': 'Skins', '보급품': 'Supplies', '보스': 'Boss', '광산': 'Mine',
    '포기': 'Give Up', '던전을 포기할까요?\\n(사용한 티켓은 돌아오지 않습니다)': 'Give up the dungeon?\\n(Used tickets are not refunded)',
    '눌러서 강화!': 'Tap to upgrade!', '절전 모드\\n\\n화면을 누르면 돌아옵니다': 'Eco Mode\\n\\nTap the screen to return',
    '남은 좀비 {0}': 'Zombies left {0}', '보스! {0}초': 'BOSS! {0}s', '반복 사냥 중 — 강화 후 보스 도전!': 'Farming — upgrade, then fight the boss!',
    '{0}초': '{0}s',
    # ---------- 진행 알림 ----------
    '보스 실패! 강화 후 다시 도전하세요': 'Boss failed! Upgrade and try again', '아직 보스를 만날 때가 아닙니다': "It's not time to face the boss yet",
    '보스 등장!': 'Boss incoming!', '아직 열리지 않았습니다': 'Not unlocked yet', '인벤토리가 가득 찼습니다': 'Inventory is full',
    '최고 등급입니다': 'Already the highest grade', '총알을 장착한 칸만 강화할 수 있습니다': 'Only slots with ammo equipped can be enhanced',
    '최고 강화입니다': 'Already max enhancement', '{0} 등급 총알은 +{1} 까지 — 더 높은 등급을 장착하세요': '{0} grade ammo caps at +{1} — equip a higher grade',
    '골드가 부족합니다': 'Not enough gold', '다이아가 부족합니다': 'Not enough gems', '이미 가진 보급품입니다': 'You already own this supply',
    '챕터 {0} 에서 열립니다': 'Unlocks at chapter {0}', '최소 1개의 총알을 장착해야 합니다': 'You must equip at least 1 ammo',
    '빈 칸에만 놓을 수 있습니다': 'Can only be placed in an empty slot', '{0}초 뒤에 다시 볼 수 있습니다': 'Available again in {0}s',
    '모든 단계를 끝냈습니다': 'All tiers completed', '아직 달성하지 못했습니다': 'Not achieved yet', '오늘은 다 받았습니다 (내일 다시)': 'All claimed today (come back tomorrow)',
    '다음 챕터에서 새 도전이 열립니다': 'A new challenge opens next chapter', '도전 {0} — 모두 처치!': 'Challenge {0} — defeat them all!',
    '오늘 티켓을 다 썼습니다': 'No tickets left today', '다이아 광산': 'Gem Mine', '보스 던전 Lv.{0}': 'Boss Dungeon Lv.{0}',
    '도전 {0} 성공!': 'Challenge {0} cleared!', '시간 초과': "Time's up", '채굴 끝!': 'Mining done!', '보스 처치!': 'Boss defeated!',
    '[테스트] 광고 대신 바로 지급': '[Test] Reward given without an ad', '결제는 준비 중입니다': 'Purchases are coming soon',
    # ---------- 창 공통 · 던전 ----------
    '확인': 'OK', '시작': 'Start', '{0}초 안에 좀비 {1}마리를 모두 처치!': 'Defeat all {1} zombies within {0}s!',
    '화면을 눌러 분대를 직접 움직이면 빨라요': 'Tap the screen to move your squad — it\'s faster', '보상: 다이아 {0} (챕터당 1번)': 'Reward: {0} gems (once per chapter)',
    '보상 없음': 'No reward', '다이아 광산 {0}단계': 'Gem Mine Tier {0}', '모든 공격이 1 피해로 고정!': 'Every hit deals exactly 1 damage!',
    '오늘 티켓': "Today's tickets", '{0}초 안에 {1}를 처치!': 'Defeat {1} within {0}s!', '체력': 'HP', '보상: 다이아 {0}': 'Reward: {0} gems',
    '받기': 'Claim', '광고 보고': 'Watch ad', '돌아가기': 'Back',
    # ---------- 스킨 · 상점 · 보급품 · 미션 · 설정 ----------
    '그림 준비 중': 'Art coming soon', '다이아 {0} 으로 살까요?': 'Buy for {0} gems?', '보유 효과': 'Owned bonus',
    '공격력 +{0}% · 공격속도 +{1}% · 돈 +{2}%': 'ATK +{0}% · SPD +{1}% · Gold +{2}%', '보유 효과 합계': 'Total owned bonus',
    '공격력 +{0}% · 공격속도 +{1}% · 돈 획득량 +{2}%': 'ATK +{0}% · SPD +{1}% · Gold +{2}%', '입는 중': 'Wearing', '보유': 'Owned',
    '스킨 구매': 'Buy Skin', '광고': 'Ad', '{0} 으로 살까요?': 'Buy for {0}?', '구매': 'Buy',
    '공격 +{0}% · 속도 +{1}% · 돈 +{2}% · 보스 +{3}%': 'ATK +{0}% · SPD +{1}% · Gold +{2}% · Boss +{3}%',
    '보유 효과 합계 ({0}/{1})': 'Total owned bonus ({0}/{1})', '공격력 +{0}% · 공격속도 +{1}% · 돈 +{2}% · 보스 +{3}%': 'ATK +{0}% · SPD +{1}% · Gold +{2}% · Boss +{3}%',
    '챕터 {0}': 'Ch. {0}', '보급품 구매': 'Buy Supply', '모든 단계 완료': 'All tiers complete', '완료': 'Done',
    '배경음': 'Music', '효과음': 'Sound', '진동': 'Vibration', '던전에서는 바꿀 수 없습니다': "Can't change during a dungeon", '저장': 'Save',
    '저장했습니다 (기기 안)': 'Saved (on this device)', '불러오기 · 구매 복원 · 문의': 'Load · Restore · Contact',
    '클라우드 저장·결제 복원·문의는 준비 중입니다': 'Cloud save, restore and contact are coming soon', '켜짐': 'On', '꺼짐': 'Off',
    '설정': 'Settings',
    # ---------- 총알 창 ----------
    '총알': 'Ammo', '리더': 'Leader', '동료 1': 'Ally 1', '동료 2': 'Ally 2', '장착': 'Equip', '강화': 'Enhance', 'D 등급 총알 획득!': 'Got D grade ammo!',
    '같은 등급 두 개를 끌어다 겹치면 합성!\\n장착 칸을 누르면 슬롯 강화 · 남는 총알은 분해 → 다이아': 'Drag two of the same grade together to merge!\\nTap an equip slot to enhance · Salvage spares → gems',
    '뽑기': 'Draw', '총알을 누르면 정보가 나옵니다.\\n장착 칸으로 끌어다 놓으면 장착!': 'Tap ammo to see its info.\\nDrag it onto an equip slot to equip!',
    '{0}발': '{0} pellets', '{0}마리 관통': 'pierces {0}', '범위 {0}': 'radius {0}', '일직선 전부': 'hits the whole line',
    '공격력 {0} · 연사 {1}초': 'ATK {0} · fires every {1}s', '넘친 공속 → 공격 x{0}': 'Excess speed → ATK x{0}',
    '(칸 +{0}, 이 등급은 +{1}까지)': '(slot +{0}, this grade caps at +{1})', '최대 +{0}': 'max +{0}', '슬롯 강화 +{0}': 'Slot +{0}',
    '공격 +{0}% · 속도 +{1}% · 보스 +{2}%': 'ATK +{0}% · SPD +{1}% · Boss +{2}%', '분해': 'Salvage', '해제': 'Unequip',
    '{0} 을(를) 분해할까요?': 'Salvage {0}?', '슬롯 강화': 'Slot Enhance', '성공 확률 {0}%': 'Success rate {0}%',
    '공격속도 +{0}% · 보스 피해 +{1}%': 'SPD +{0}% · Boss DMG +{1}%', '강화 성공!': 'Enhanced!', '강화 실패… (재료만 사라짐)': 'Enhance failed… (materials lost)',
    '합성!': 'Merged!', '장착 중인 총알은 해제한 뒤 합성할 수 있습니다': 'Unequip ammo before merging it',
    '단발': 'Single', '산탄': 'Shotgun', '속사': 'Rapid', '관통': 'Pierce', '레이저': 'Laser', '미사일': 'Missile',
    # ---------- 표: 스킨 ----------
    '기본 카피': 'Classic Capy', '밀짚모자': 'Straw Hat', '천사 날개': 'Angel Wings', '새싹': 'Sprout', '꿀꿀 돼지': 'Piggy', '악어 잠옷': 'Croc Pajamas',
    '원시인': 'Caveman', '검은 두건': 'Black Hood', '그림자 망토': 'Shadow Cloak', '공룡 잠옷': 'Dino Pajamas', '모험가 배낭': "Explorer's Pack",
    '박쥐 날개': 'Bat Wings', '멜빵바지': 'Overalls', '꼬꼬 잠옷': 'Chicken Pajamas', '빨간 망토': 'Red Cape', '상어 잠옷': 'Shark Pajamas',
    '줄무늬 고양이': 'Tabby Cat', '해적 선장': 'Pirate Captain',
    # ---------- 표: 총알 ----------
    '리볼버': 'Revolver', '소드오프 샷건': 'Sawed-off Shotgun', '우지': 'Uzi', '상위 리볼버': 'Elite Revolver', '펌프액션 샷건': 'Pump Shotgun',
    '돌격소총': 'Assault Rifle', '매그넘': 'Magnum', '전투 샷건': 'Combat Shotgun', '중기관총': 'Heavy MG', '관통 라이플': 'Piercing Rifle',
    '황금 매그넘': 'Golden Magnum', '드래곤 샷건': 'Dragon Shotgun', '미니건': 'Minigun', '레일 라이플': 'Rail Rifle', '레이저 건': 'Laser Gun',
    '미사일 런처': 'Missile Launcher', '플라즈마 레이저': 'Plasma Laser', '다연장 미사일': 'Multi Missile', '오메가 레이저': 'Omega Laser',
    '핵 미사일': 'Nuke Missile', '하이퍼 레이저': 'Hyper Laser', '전술 핵 미사일': 'Tactical Nuke', '퀘이사 레이저': 'Quasar Laser',
    '반물질 미사일': 'Antimatter Missile', '네뷸라 레이저': 'Nebula Laser', '중력자 미사일': 'Graviton Missile', '노바 레이저': 'Nova Laser',
    '초신성 미사일': 'Supernova Missile', '보이드 레이저': 'Void Laser', '차원파괴 미사일': 'Rift Missile', '제네시스 레이저': 'Genesis Laser',
    '아포칼립스 미사일': 'Apocalypse Missile',
    # ---------- 표: 좀비 ----------
    '기어오는 좀비': 'Crawler', '칼 좀비': 'Knife Zombie', '좀비': 'Zombie', '모자 좀비': 'Hat Zombie', '양동이 좀비': 'Bucket Zombie',
    '갑옷 좀비': 'Armored Zombie', '방망이 좀비': 'Bat Zombie', '농부 좀비': 'Farmer Zombie', '미라': 'Mummy', '뚱보 좀비': 'Fat Zombie',
    '사신 좀비': 'Reaper Zombie', '변이 오우거': 'Mutant Ogre',
    # ---------- 표: 보급품 ----------
    '비상 건빵': 'Emergency Crackers', '통조림 캔': 'Canned Food', '정수 필터': 'Water Filter', '야전 반합': 'Mess Kit', '에너지 바': 'Energy Bar',
    '구급 상자': 'First Aid Kit', '아드레날린 주사': 'Adrenaline Shot', '군용 무전기': 'Field Radio', '야간 투시경': 'Night Vision',
    '방탄 조끼': 'Bulletproof Vest', '전술 탄띠': 'Tactical Bandolier', '고농축 스팀팩': 'Mega Stimpack', '열화우라늄 탄심': 'DU Core',
    '냉각 모듈': 'Cooling Module', '플라즈마 배터리': 'Plasma Battery', '양자 증폭기': 'Quantum Amplifier', '중력 제어 코어': 'Gravity Core',
    '나노 수복 혈청': 'Nano Serum', '바이오 리액터': 'Bio Reactor', '초전도 코일': 'Superconductor Coil', '암흑물질 셀': 'Dark Matter Cell',
    '차원 도약 엔진': 'Warp Engine', '항성 에너지 핵': 'Stellar Core', '영구기관 모듈': 'Perpetual Engine', '카피바라의 성배': 'Capybara Grail',
    '일반': 'Common', '고급': 'Uncommon', '희귀': 'Rare', '영웅': 'Epic', '전설': 'Legendary', '신화': 'Mythic', '초월': 'Transcendent',
    '불멸': 'Immortal', '태초': 'Primordial',
    # ---------- 표: 미션 · 상점 · 챕터 · 열림 · 광산 ----------
    '좀비 처치하기': 'Defeat zombies', '스테이지 클리어하기': 'Clear stages', '광산 플레이하기': 'Play the mine', '보스 처치하기': 'Defeat bosses',
    '총알 합성하기': 'Merge ammo',
    '공격력 x2 (5분)': 'Attack x2 (5 min)', '돈 획득량 x2 (5분)': 'Gold x2 (5 min)', '강제 광고 제거': 'Remove Forced Ads', '모든 광고 제거': 'Remove All Ads',
    '자동강화 패키지': 'Auto Upgrade Pack', '영구 3배 버프': 'Permanent x3 Buff', '광산 티켓 확장': 'Mine Ticket Pack', '핫딜 레이저 총알': 'Hot Deal: Laser Ammo',
    '다이아 한 줌 (500)': 'Handful of Gems (500)', '다이아 주머니 (1,800)': 'Pouch of Gems (1,800)', '다이아 상자 (4,200)': 'Chest of Gems (4,200)',
    '30분 방치 골드 상자': '30-min Gold Chest', '무료 다이아': 'Free Gems',
    '광고 1편 · 연속 시청 시 시간 누적': '1 ad · time stacks when watched again', '자동 전면 광고 제거 + 돼지저금통': 'No more forced ads + piggy bank',
    '보상형 광고도 바로 보상': 'Rewards without watching ads', '하단 강화 자동 + 다이아 500': 'Upgrades buy themselves + 500 gems',
    '공격력·돈 x3 영구 + 다이아 1,000': 'Permanent x3 attack & gold + 1,000 gems', '하루 광산 티켓 2 → 6 + 다이아 500': 'Daily mine tickets 2 → 6 + 500 gems',
    'S 등급 레이저 건 즉시 지급': 'Get an S grade Laser Gun now', '다이아 500개 즉시 획득': 'Get 500 gems now', '다이아 1,800개 즉시 획득': 'Get 1,800 gems now',
    '다이아 4,200개 즉시 획득': 'Get 4,200 gems now', '현재 스테이지 기준 600마리 처치 골드': 'Gold of 600 zombies at your current stage',
    '광고 1편 · 다이아 40 · 하루 3번': '1 ad · 40 gems · 3 times a day',
    '폐허 숲길': 'Ruined Forest Trail', '오염된 해안': 'Polluted Coast', '사막 도로': 'Desert Road', '얼어붙은 설원': 'Frozen Snowfield',
    '지하 연구소': 'Underground Lab', '오염 구역': 'Contaminated Zone',
    '무기는 다이아 100 또는 도전 클리어 에서 열립니다': 'Weapons unlock with 100 gems or a challenge clear',
    'Stage 1-2 에서 열립니다': 'Unlocks at Stage 1-2', 'Stage 4-1 에서 열립니다': 'Unlocks at Stage 4-1', 'Stage 6-1 에서 열립니다': 'Unlocks at Stage 6-1',
    'Stage 10-1 에서 열립니다': 'Unlocks at Stage 10-1', '무기 메뉴가 열렸습니다!': 'Weapons unlocked!', '스킨이 열렸습니다!': 'Skins unlocked!',
    '다이아 광산이 열렸습니다!': 'Gem Mine unlocked!', '보급품이 열렸습니다!': 'Supplies unlocked!', '보스 던전이 열렸습니다!': 'Boss Dungeon unlocked!',
    '큰 푸른 보석': 'Big Blue Gem', '작은 푸른 보석': 'Small Blue Gem', '회색 바위': 'Gray Rock', '거대 황금 크리스탈': 'Huge Gold Crystal',
}

# 10-24 총알 이름 규칙 "종류_등급" (Weapon_Master Name) → 영어 "Type_Grade"
for _ko, _en in (('기본탄', 'Basic Round'), ('산탄', 'Buckshot'), ('속사탄', 'Rapid Round'), ('관통탄', 'Piercing Round'), ('레이저탄', 'Laser Round'), ('미사일탄', 'Missile')):
    for _g in ('D', 'C', 'B', 'A', 'S', 'SS', 'SSS', 'R', 'SR', 'SSR', 'U', 'SU', 'SSU'):
        EN[f'{_ko}_{_g}'] = f'{_en}_{_g}'
EN['레이저탄_S 즉시 지급'] = 'Get a Laser Round_S now'


def code_strings():
    """(글자, 파일) — 주석·Debug 줄 제외, $"" 안 한글은 경고"""
    out, warn = [], []
    # 따옴표 사이 조각을 겹쳐서 모두 봄: $"{L.T("광산")} …" 안의 "광산" 도 잡힘.
    # (따옴표 밖 조각은 코드라 한글이 없음 → 한글 있는 조각 = 글자)
    seg = re.compile(r'(?="((?:[^"\\\n]|\\.)*)")')
    for f in sorted(glob.glob(os.path.join(SCRIPTS, '**', '*.cs'), recursive=True)):
        rel = os.path.relpath(f, SCRIPTS).replace('\\', '/')
        if rel.startswith('Editor/'): continue
        for n, line in enumerate(io.open(f, encoding='utf-8'), 1):
            s = line.strip()
            if 'Debug.Log' in line or s.startswith('//'): continue
            # 줄 끝 주석 떼기 (글자 밖의 // 만)
            q, cut = False, len(line)
            for i, ch in enumerate(line):
                if ch == '"' and (i == 0 or line[i - 1] != '\\'): q = not q
                if not q and line.startswith('//', i): cut = i; break
            code = line[:cut]
            for m in seg.finditer(code):
                t = m.group(1)
                if not KO.search(t) or t in SKIP or '{L.' in t or 'L.T(' in t: continue
                # $"" 안에 한글이 그대로 박힌 글자 = 번역 안 됨 → 알림
                if re.search(r'\{[^}]*\}', t) and not re.fullmatch(r'[^{}]*(\{\d\}[^{}]*)*', t): warn.append(f'{rel}:{n}: {t}'); continue
                out.append((t, rel))
    return out, warn


def table_strings():
    out = []
    for name, cols in TABLE_COLS:
        p = os.path.join(TABLES, name + '.csv')
        if not os.path.exists(p): continue
        for r in csv.DictReader(io.open(p, encoding='utf-8-sig')):
            for c in cols:
                v = (r.get(c) or '').strip()
                if v and v != '-' and KO.search(v): out.append((v.replace('\n', '\\n'), name + '.' + c))
    return out


def main():
    old = {}
    if os.path.exists(OUT):
        ws = openpyxl.load_workbook(OUT).worksheets[0]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]: old[str(row[0])] = (row[1] or '', row[2] or '')
    code, warn = code_strings()
    rows, seen = [], set()
    for t, where in code + table_strings():
        if t in seen: continue
        seen.add(t)
        en = old[t][0] if t in old and old[t][0] else EN.get(t, '')
        rows.append((t, en, where))
    # 코드·표에서 사라진 줄도 지우지 않고 맨 아래에 둠 (직접 쓴 영어 보존)
    gone = [(k, v[0], '(지금은 안 씀) ' + v[1].replace('(지금은 안 씀) ', '')) for k, v in old.items() if k not in seen]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'String'
    ws.append(['Key', 'EN', 'Where'])
    for r in rows + gone: ws.append(list(r))
    head = PatternFill('solid', fgColor='2F5233')
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = head
    miss = PatternFill('solid', fgColor='FFE699')
    for row in ws.iter_rows(min_row=2):
        if not row[1].value: row[1].fill = miss          # 영어가 빈 칸 = 노란색
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 60; ws.column_dimensions['B'].width = 60; ws.column_dimensions['C'].width = 30
    ws.freeze_panes = 'A2'
    wb.save(OUT)
    empty = [r[0] for r in rows if not r[1]]
    print(f'String.xlsx: {len(rows)}줄 (영어 빈 칸 {len(empty)}) · 안 쓰는 줄 {len(gone)}')
    for e in empty: print('  영어 없음:', e)
    for w in warn: print('  [경고] 번역 못 하는 글자($"" 안 한글):', w)
    # String.xlsx 만 csv 로 (다른 표는 건드리지 않음)
    import xlsx2csv
    wb = openpyxl.load_workbook(OUT, data_only=True)
    text = xlsx2csv.sheet_csv(wb.worksheets[0])
    with io.open(os.path.join(TABLES, 'String.csv'), 'w', encoding='utf-8-sig', newline='') as fp: fp.write(text)
    print('String.csv 씀')


if __name__ == '__main__':
    main()
