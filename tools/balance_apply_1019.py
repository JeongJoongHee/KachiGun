# -*- coding: utf-8 -*-
"""10-19 밸런스 적용: Player.xlsx Gold_Upgrade 다시 만들기 + 광고 배율 x2 (Buff·Config·Shop·Unlock)"""
import sys, math, json, io
import openpyxl
sys.path.insert(0, r'D:/00 Final_Game/tools')
import balance_tune as T

TB = r'D:/00 Final_Game/Kachi_Gun/Table/'
K = [0.47, 0.47, 0.47, 0.47, 0.45]   # 구간별 (6달: 하루 4시간·무과금)
BOUNDS = (80, 250, 500, 1000)
pars, rows = T.candidate(K, spd_mul=True, bounds=BOUNDS)


def fmt(v):
    """게임 Num.F 와 같은 표기 (내림)"""
    if v < 1000: return str(int(math.floor(v)))
    k = int(math.floor(math.log10(v) / 3)); m = v / 10 ** (k * 3)
    if m >= 1000: k += 1; m /= 1000
    named = ['', 'k', 'm', 'b', 't']
    unit = named[k] if k < 5 else chr(ord('a') + (k - 5) // 26 % 26) + chr(ord('a') + (k - 5) % 26)
    tr = lambda x, d: math.floor(x * 10 ** d) / 10 ** d
    s = ('%.2f' % tr(m, 2)) if m < 10 else ('%.1f' % tr(m, 1)) if m < 100 else str(int(math.floor(m)))
    if '.' in s: s = s.rstrip('0').rstrip('.')
    return s + unit


# ---------- Player.xlsx / Gold_Upgrade_1500 ----------
wb = openpyxl.load_workbook(TB + 'Player.xlsx')
ws = wb['Gold_Upgrade_1500']
head = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(head)}
assert 'Spd_Add_Per_Lv' in col
ws.cell(1, col['Spd_Add_Per_Lv']).value = 'Spd_Growth_Rate'
col['Spd_Growth_Rate'] = col.pop('Spd_Add_Per_Lv')
cum = 0
for lv, r in enumerate(rows):
    R = lv + 2
    cum += r['Upgrade_Cost']
    vals = {
        'Lv': lv, 'Target_Chapter': min(150, lv // 10 + 1), 'Cost_Growth_Rate': round(r['Cost_Growth_Rate'], 5),
        'Upgrade_Cost': round(r['Upgrade_Cost'], 2), 'Cost_표기': fmt(r['Upgrade_Cost']), 'Cum_Cost_표기': fmt(cum),
        'Stat_Growth_Rate': round(r['Stat_Growth_Rate'], 5), 'Atk_Power (%)': round(r['Atk_Power'], 4), 'Atk_Mul (배율)': round(r['Atk_Power'] / 100, 6),
        'Atk_표기': fmt(r['Atk_Power']), 'Spd_Growth_Rate': round(r['Stat_Growth_Rate'], 5), 'Atk_Speed (%)': round(r['Atk_Speed'], 4),
        'Spd_Mul (배율)': round(r['Atk_Speed'] / 100, 6), 'Gold_Growth_Rate': round(r['Gold_Growth_Rate'], 5),
        'Money_Gain (%)': round(r['Money_Gain'], 4), 'Gold_Mul (배율)': round(r['Money_Gain'] / 100, 6),
    }
    for k, v in vals.items(): ws.cell(R, col[k]).value = v
wb.save(TB + 'Player.xlsx')

# ---------- Config.xlsx: 구간 설명 · 광고 배율 ----------
wb = openpyxl.load_workbook(TB + 'Config.xlsx')
ws = wb['Bracket_Growth_Rules']
head = [c.value for c in ws[1]]
c = {h: i + 1 for i, h in enumerate(head)}
ws.cell(1, c['Upgrade_Spd_Add']).value = 'Upgrade_Spd_Growth'
lvr = ['Lv.0 ~ 80', 'Lv.81 ~ 250', 'Lv.251 ~ 500', 'Lv.501 ~ 1000', 'Lv.1001 ~ 1500']
for i in range(5):
    R = i + 2
    cg, a, m = pars[i]
    ws.cell(R, c['Player_Lv_Range']).value = lvr[i]
    ws.cell(R, c['Upgrade_Cost_Growth']).value = round(cg, 5)
    ws.cell(R, c['Upgrade_Stat_Growth']).value = round(a, 5)
    ws.cell(R, c['Upgrade_Spd_Add']).value = round(a, 5)
ws = wb['Global_Config']
for row in ws.iter_rows(min_row=2):
    if row[0].value in ('BuffMul', 'AdMultiplier'): row[1].value = 2
wb.save(TB + 'Config.xlsx')


def x2(s):
    return s.replace('x3', 'x2').replace('3배', '2배') if isinstance(s, str) else s


# ---------- Buff.xlsx ----------
wb = openpyxl.load_workbook(TB + 'Buff.xlsx')
ws = wb['Buff_Table']
head = [c.value for c in ws[1]]
c = {h: i + 1 for i, h in enumerate(head)}
for row in ws.iter_rows(min_row=2):
    if row[0].value in ('buff_atk', 'buff_gold', 'dungeon_ad_x3'):
        row[c['Multiplier'] - 1].value = 2
        row[c['Name'] - 1].value = x2(row[c['Name'] - 1].value)
        row[c['Description'] - 1].value = x2(row[c['Description'] - 1].value)
wb.save(TB + 'Buff.xlsx')

# ---------- Shop.xlsx (광고 버프 두 줄) ----------
wb = openpyxl.load_workbook(TB + 'Shop.xlsx')
ws = wb['Shop_Products']
for row in ws.iter_rows(min_row=2):
    if row[1].value in ('buff_atk', 'buff_gold'):
        for cell in row: cell.value = x2(cell.value)
wb.save(TB + 'Shop.xlsx')

# ---------- Unlock.xlsx (버프 이름) ----------
wb = openpyxl.load_workbook(TB + 'Unlock.xlsx')
ws = wb['Unlock_Roadmap']
for row in ws.iter_rows(min_row=2):
    if row[0].value == 'buff':
        for cell in row: cell.value = x2(cell.value)
wb.save(TB + 'Unlock.xlsx')
print('ok', [tuple(round(x, 5) for x in p) for p in pars])
